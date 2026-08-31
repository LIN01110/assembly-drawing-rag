"""多输入适配层（ingest）：把三类输入 + 人工补全统一收敛到标准零件 JSON。

入口：
  ingest_ocr_result(path)  — integrate.py 产出的 result.json（图片/PDF 扫描件经 YOLO+OCR 后的结构化输出）
  ingest_pdf(path)         — 文字型 PDF 抽文本层；扫描件返回 needs_ocr=True（应转图后进 OCR 管线）
  ingest_docx(path)        — Word 检验单/技术协议：抽段落+表格
  merge_manual(part, manual, actor) — 人工补全合并：人工值 > OCR 值，逐字段留痕
  所有入口产出统一 part dict，附带 _source/_ingest_meta 供审计。
"""
import json
import re
from pathlib import Path

# OCR structured_info 字段 → part schema 的映射
FIELD_MAP = {
    "材料": "material",
    "名称": "part_type",
    "图号": "drawing_no",
    "比例": "scale",
    "重量": "weight",
}


def _extract_tech_blob(text: str) -> str:
    """从 OCR 全文/技术要求文本中提取特征关键词。"""
    kws = []
    for kw in ("调质", "淬火", "时效", "固溶", "渗碳", "氮化", "同轴度", "垂直度", "平行度",
               "螺纹", "键槽", "Ra", "±"):
        if kw in text:
            kws.append(kw)
    return " ".join(kws)


def ingest_ocr_result(path: str | Path) -> dict:
    """integrate.py 的 result.json → part schema。"""
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    info = data.get("structured_info", {})
    fields = info.get("fields", info) if isinstance(info, dict) else {}
    part: dict = {"_source": "ocr", "_ingest_meta": {"file": str(path)}}
    for cn, en in FIELD_MAP.items():
        if fields.get(cn):
            part[en] = fields[cn]
    # 技术要求与全文
    all_text = " ".join(t.get("text", "") for t in data.get("all_text", []))
    tech = info.get("tech_requirements", "") or all_text
    part["tech_requirements"] = tech[:500]
    part["features"] = _extract_tech_blob(tech + " " + all_text)
    part.setdefault("part_type", "未识别零件")
    return part


def ingest_pdf(path: str | Path) -> dict:
    """文字型 PDF 抽文本层；无文字层（扫描件）标记 needs_ocr。"""
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    text = "\n".join(page.extract_text() or "" for page in reader.pages).strip()
    if len(text) < 20:
        return {"_source": "pdf", "needs_ocr": True,
                "_ingest_meta": {"file": str(path), "reason": "无文字层，疑似扫描件，应转图走 OCR 管线"}}
    part: dict = {"_source": "pdf", "_ingest_meta": {"file": str(path), "pages": len(reader.pages)}}
    part["tech_requirements"] = text[:500]
    part["features"] = _extract_tech_blob(text)
    m = re.search(r"材料[:：]\s*([0-9A-Za-z一-龥#-]+)", text)
    if m:
        part["material"] = m.group(1)
    part.setdefault("part_type", "PDF文档零件")
    return part


def ingest_docx(path: str | Path) -> dict:
    """Word 检验单/技术协议：段落 + 表格抽取，'键: 值' 对入字段。"""
    import docx
    doc = docx.Document(str(path))
    part: dict = {"_source": "docx", "_ingest_meta": {"file": str(path)}}
    texts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if len(cells) >= 2 and cells[0]:
                texts.append(f"{cells[0]}: {cells[1]}")
                for cn, en in FIELD_MAP.items():
                    if cn in cells[0] and cells[1]:
                        part[en] = cells[1]
    blob = "\n".join(texts)
    part["tech_requirements"] = blob[:500]
    part["features"] = _extract_tech_blob(blob)
    part.setdefault("part_type", "检验单零件")
    return part


def merge_manual(part: dict, manual: dict, actor: str) -> tuple[dict, list[str]]:
    """人工补全合并：人工值 > OCR/自动值，逐字段留痕。"""
    out = dict(part)
    changes = []
    for k, v in manual.items():
        if v in (None, ""):
            continue
        old = out.get(k)
        if old != v:
            changes.append(f"{k}: {old!r} → {v!r}（人工补全 by {actor}）")
            out[k] = v
    out["_manual_by"] = actor
    return out, changes
