"""多输入适配层端到端测试：OCR result.json / PDF / DOCX / 人工补全 → graph 全流程。

运行：python scripts/test_ingest.py
产物：自动构造 test_assets/sample.pdf 与 sample.docx（文字型），OCR 用真实产物。
"""
import json
import os
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

os.environ["PROCESS_RAG_NO_LLM"] = "1"  # 测试确定性，LLM 路径已由 run_demo 验证

from langgraph.types import Command

from src import db
from src.graph import build_graph
from src.ingest import ingest_docx, ingest_ocr_result, ingest_pdf, merge_manual

ASSETS = ROOT / "scripts" / "test_assets"
OCR_RESULT = ROOT.parents[0] / "yolo" / "predict" / "voting_fix_test" / "big_scan0_result.json"


def make_pdf(path: Path) -> None:
    """用 reportlab 造一份文字型 PDF（CID 中文字体）。"""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfgen import canvas
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    c = canvas.Canvas(str(path))
    c.setFont("STSong-Light", 12)
    c.drawString(72, 760, "技术协议：传动轴")
    c.drawString(72, 740, "材料: 40Cr，调质 HRC28-32")
    c.drawString(72, 720, "外圆公差 ±0.02，Ra 1.6，同轴度 0.02")
    c.save()


def make_docx(path: Path) -> None:
    import docx
    doc = docx.Document()
    doc.add_paragraph("零件检验单")
    table = doc.add_table(rows=3, cols=2)
    table.rows[0].cells[0].text = "材料"; table.rows[0].cells[1].text = "45"
    table.rows[1].cells[0].text = "名称"; table.rows[1].cells[1].text = "联轴器半体"
    table.rows[2].cells[0].text = "图号"; table.rows[2].cells[1].text = "LT-2026-081"
    doc.add_paragraph("技术要求：调质处理，Ra 3.2，内孔 H7")
    doc.save(str(path))


def run_graph(part: dict, tag: str) -> dict:
    g = build_graph()
    cfg = {"configurable": {"thread_id": f"ingest-{tag}"}}
    r = g.invoke({"part": part}, cfg)
    while "__interrupt__" in r:
        r = g.invoke(Command(resume={"action": "approve", "actor": "测试", "role": "批准人"}), cfg)
    return r["final"]


def main() -> None:
    db.init_db()
    ASSETS.mkdir(exist_ok=True, parents=True)
    pdf_path, docx_path = ASSETS / "sample.pdf", ASSETS / "sample.docx"
    make_pdf(pdf_path)
    make_docx(docx_path)

    print("== 1. OCR result.json（真实产物 big_scan0）==")
    part1 = ingest_ocr_result(OCR_RESULT)
    print("  抽取:", {k: v for k, v in part1.items() if not k.startswith("_")})
    # 人工补全：OCR 材料 Q45 存疑，人工确认 45 钢
    part1, changes = merge_manual(part1, {"material": "45", "part_type": "盖板", "tolerance_it": 9}, actor="检验员甲")
    print("  人工补全留痕:", changes)
    f1 = run_graph(part1, "ocr")
    print("  工艺:", f1["steps"], "| 车间:", f1.get("workshop"), "| 归档:", f1.get("archived"))

    print("== 2. 文字型 PDF ==")
    part2 = ingest_pdf(pdf_path)
    print("  抽取:", {k: v for k, v in part2.items() if not k.startswith("_") and k != "tech_requirements"})
    f2 = run_graph(part2, "pdf")
    print("  工艺:", f2["steps"], "| 车间:", f2.get("workshop"))

    print("== 3. Word 检验单 ==")
    part3 = ingest_docx(docx_path)
    print("  抽取:", {k: v for k, v in part3.items() if not k.startswith("_") and k != "tech_requirements"})
    f3 = run_graph(part3, "docx")
    print("  工艺:", f3["steps"], "| 车间:", f3.get("workshop"))

    print("== 4. 扫描件 PDF 路径（needs_ocr 标记）==")
    empty_pdf = ASSETS / "empty.pdf"
    from pypdf import PdfWriter
    w = PdfWriter(); w.add_blank_page(200, 200)
    with open(empty_pdf, "wb") as f:
        w.write(f)
    print(" ", ingest_pdf(empty_pdf))


if __name__ == "__main__":
    main()
