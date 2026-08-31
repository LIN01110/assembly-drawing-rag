"""
build_gt.py — 把人工校正后的 Excel 转成 OCR 真值 JSON

用法（仓库根目录）：
  .conda/python.exe ocr_truth/build_gt.py
  .conda/python.exe ocr_truth/build_gt.py --xlsx ocr_truth/ocr_truth_draft.xlsx

产出：ocr_truth/ground_truth.json
  { "<图片stem>": {"fields": {"材料": "Q235", ...}, "tech_requirements": "1. ...\n2. ..."} }
空单元格 = 图纸无此字段，不写入 JSON。
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
OUT_JSON = ROOT / "ocr_truth" / "ground_truth.json"

FIELDS = ["物品名称", "材料", "图号", "数量", "张数", "重量",
          "比例", "设计", "审核", "工艺", "标准化", "批准"]


def main() -> None:
    import openpyxl

    parser = argparse.ArgumentParser()
    parser.add_argument("--xlsx", default=str(ROOT / "ocr_truth" / "ocr_truth_draft.xlsx"))
    args = parser.parse_args()

    wb = openpyxl.load_workbook(args.xlsx)
    ws = wb["字段校正"]
    headers = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(headers)}

    gt: dict[str, dict] = {}
    n_cells = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        stem = row[col["图片"]]
        if not stem:
            continue
        fields = {}
        for f in FIELDS:
            v = row[col[f]]
            if v is not None and str(v).strip():
                fields[f] = str(v).strip()
                n_cells += 1
        tech = row[col["技术要求(草稿)"]]
        gt[str(stem)] = {
            "fields": fields,
            "tech_requirements": str(tech).strip() if tech else "",
        }

    OUT_JSON.write_text(json.dumps(gt, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"真值已导出: {OUT_JSON}")
    print(f"共 {len(gt)} 张图，{n_cells} 个非空字段。")
    print("下一步：用真值跑 OCR 臂评估（字段准确率 / CER），回填简历数字。")


if __name__ == "__main__":
    main()
