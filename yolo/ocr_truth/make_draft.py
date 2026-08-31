"""
make_draft.py — OCR 真值草稿生成器

读取 predict/gt_batch/*_result.json（integrate.py 批处理输出），
生成预填好 OCR 草稿的 Excel，人工只需对照图纸改错字。

用法（仓库根目录）：
  .conda/python.exe ocr_truth/make_draft.py

产出：
  ocr_truth/ocr_truth_draft.xlsx
    - Sheet「字段校正」：每行一张图，每列一个字段，单元格已预填草稿，**直接改错字即可**；
      图纸里没有的字段把单元格清空。
    - Sheet「原始OCR文本」：每张图的全部 OCR 原文，校正时参考用。
  ocr_truth/vis/*.jpg：检测可视化图（对照图纸时用）。
"""

from __future__ import annotations

import json
import shutil
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
BATCH_DIR = ROOT / "predict" / "gt_batch"
OUT_DIR = ROOT / "ocr_truth"
OUT_XLSX = OUT_DIR / "ocr_truth_draft.xlsx"
VIS_DIR = OUT_DIR / "vis"

# 展示名 -> 表格信息里的候选键（与 integrate.py save_mechanical_drawing_txt 的映射一致）
FIELD_MAPPING = [
    ("物品名称", ["物品名称", "物品"]),
    ("材料", ["材料", "材质", "牌号"]),
    ("图号", ["图号", "图纸编号"]),
    ("数量", ["数量", "件数"]),
    ("张数", ["张数"]),
    ("重量", ["重量", "质量"]),
    ("比例", ["比例", "SCALE"]),
    ("设计", ["设计"]),
    ("审核", ["审核"]),
    ("工艺", ["工艺"]),
    ("标准化", ["标准化"]),
    ("批准", ["批准"]),
]


def pick_field(table_info: dict, keys: list[str]) -> str:
    for k in keys:
        if k in table_info and str(table_info[k]).strip():
            return str(table_info[k]).strip()
    return ""


def main() -> None:
    import openpyxl
    from openpyxl.styles import Alignment, Font

    result_files = sorted(BATCH_DIR.glob("*_result.json"))
    if not result_files:
        raise SystemExit(f"未找到批处理结果: {BATCH_DIR}（先跑 integrate.py --mode batch）")

    OUT_DIR.mkdir(exist_ok=True)
    VIS_DIR.mkdir(exist_ok=True)

    wb = openpyxl.Workbook()

    # --- Sheet 1: 字段校正 ---
    ws = wb.active
    ws.title = "字段校正"
    headers = ["图片", "原图路径", "可视化图"] + [name for name, _ in FIELD_MAPPING] + ["技术要求(草稿)"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)

    # --- Sheet 2: 原始OCR文本 ---
    ws2 = wb.create_sheet("原始OCR文本")
    ws2.append(["图片", "原始OCR文本"])
    for c in ws2[1]:
        c.font = Font(bold=True)

    for rf in result_files:
        data = json.loads(rf.read_text(encoding="utf-8"))
        stem = Path(data.get("image", rf.stem.replace("_result", ""))).stem
        si = data.get("structured_info", {})
        table_info = si.get("表格信息", {})
        tech_items = si.get("技术要求", [])
        tech_text = "\n".join(tech_items) if tech_items else ""
        raw_text = si.get("原始OCR文本_全部", "") or si.get("技术要求_原始", "")

        # 复制可视化图
        vis_src = BATCH_DIR / f"{stem}_visualization.jpg"
        vis_rel = ""
        if vis_src.exists():
            vis_dst = VIS_DIR / vis_src.name
            shutil.copy2(vis_src, vis_dst)
            vis_rel = str(vis_dst)

        img_path = str(ROOT / "mydata" / "images" / "train" / f"{stem}.png")
        ws.append([stem, img_path, vis_rel]
                  + [pick_field(table_info, keys) for _, keys in FIELD_MAPPING]
                  + [tech_text])
        ws2.append([stem, raw_text])

    # 列宽与换行
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 40
    for col in range(4, 4 + len(FIELD_MAPPING)):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 14
    last_col = openpyxl.utils.get_column_letter(3 + len(FIELD_MAPPING) + 1)
    ws.column_dimensions[last_col].width = 50
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws2.column_dimensions["A"].width = 14
    ws2.column_dimensions["B"].width = 100
    for row in ws2.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(OUT_XLSX)
    print(f"草稿已生成: {OUT_XLSX}")
    print(f"可视化图: {VIS_DIR}（{len(list(VIS_DIR.glob('*.jpg')))} 张）")
    print(f"共 {len(result_files)} 张图纸。校正后运行: .conda/python.exe ocr_truth/build_gt.py")


if __name__ == "__main__":
    sys.exit(main())
